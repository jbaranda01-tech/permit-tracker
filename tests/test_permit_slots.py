"""End-to-end tests for permit add / toggle and company-wide (blanket) permit types.

Run from the repo root:  python3 -m pytest -q tests/test_permit_slots.py
Uses a throwaway SQLite DB in a temp dir; importing `app` runs the boot block
(create_all + migrations). The admin user is created here with a pbkdf2 hash
because some local Pythons lack hashlib.scrypt (werkzeug's default).
"""
import io
import os
import pathlib
import sys
import tempfile
from datetime import date, timedelta

REPO = pathlib.Path(__file__).resolve().parents[1]
DB = pathlib.Path(tempfile.mkdtemp(prefix='permit-tracker-test-')) / 'test_permits.db'
os.environ['SECRET_KEY'] = 'test'
os.environ['DATABASE_URL'] = 'sqlite:///' + str(DB)      # absolute path => 4 slashes
os.environ.pop('DATABASE_PRIVATE_URL', None)
os.environ['ENABLE_SCHEDULER'] = 'false'
sys.path.insert(0, str(REPO))

import pytest  # noqa: E402
import app as appmod  # noqa: E402  (boot: create_all + migrations + admin user)
from flask import url_for  # noqa: E402
from sqlalchemy.exc import IntegrityError  # noqa: E402
import models  # noqa: E402
from models import (db, Employee, EmployeePermit, Equipment, EquipmentPermit,  # noqa: E402
                    FileStorage, sync_vehicle_insurance_permit)

app = appmod.app
app.config['TESTING'] = True
app.config['SERVER_NAME'] = 'localhost'

_counter = {'n': 0}


def _uniq(prefix):
    _counter['n'] += 1
    return f'{prefix} {_counter["n"]}'


def u(endpoint, **kw):
    with app.test_request_context():
        return url_for(endpoint, **kw)


def _ensure_admin():
    # The boot block creates admin/admin123 with werkzeug's default scrypt hash,
    # which this machine's Python (LibreSSL, no hashlib.scrypt) cannot produce.
    from werkzeug.security import generate_password_hash
    from models import User
    with app.app_context():
        if not User.query.filter_by(username='admin').first():
            admin = User(username='admin', email='admin@example.com', role='admin',
                         password_hash=generate_password_hash('admin123', method='pbkdf2:sha256'))
            db.session.add(admin)
            db.session.commit()


@pytest.fixture(scope='module')
def client():
    _ensure_admin()
    c = app.test_client()
    r = c.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)
    assert r.status_code == 200
    return c


def new_emp(company='LB', archived=False):
    with app.app_context():
        emp = Employee(name=_uniq('Test Emp'), company=company)
        if archived:
            from datetime import datetime
            emp.archived_at = datetime.utcnow()
            emp.status = 'inactivo'
        db.session.add(emp)
        db.session.commit()
        return emp.id


def new_equipment(company='LB'):
    """Mirror equipment_new's default-slot loop + insurance sync."""
    with app.app_context():
        eq = Equipment(company=company, unit_number=_uniq('U'), model='Camion', equipment_class='truck')
        db.session.add(eq)
        db.session.flush()
        for code, _ in models.EQUIPMENT_PERMIT_TYPES:
            if code != 'OTHER':
                db.session.add(EquipmentPermit(equipment_id=eq.id, permit_type=code, applicability='YES'))
        sync_vehicle_insurance_permit(eq)
        db.session.commit()
        return eq.id


def emp_rows(emp_id, ptype):
    with app.app_context():
        return EmployeePermit.query.filter_by(employee_id=emp_id, permit_type=ptype) \
            .order_by(EmployeePermit.id).all()


def eq_rows(eq_id, ptype):
    with app.app_context():
        return EquipmentPermit.query.filter_by(equipment_id=eq_id, permit_type=ptype) \
            .order_by(EquipmentPermit.id).all()


def add_emp_permit(client, emp_id, ptype, name='', exp=''):
    return client.post(u('employee_permit_new', emp_id=emp_id),
                       data={'permit_type': ptype, 'permit_name': name, 'expiration_date': exp},
                       follow_redirects=True)


def add_eq_permit(client, eq_id, ptype, name='', exp=''):
    return client.post(u('equipment_permit_new', eq_id=eq_id),
                       data={'permit_type': ptype, 'permit_name': name, 'expiration_date': exp},
                       follow_redirects=True)


def toggle_emp(client, emp_id, pid, value):
    return client.post(u('employee_permit_toggle', emp_id=emp_id, permit_id=pid),
                       data={'applicability': value}, follow_redirects=True)


def toggle_eq(client, eq_id, pid, value):
    return client.post(u('equipment_permit_toggle', eq_id=eq_id, permit_id=pid),
                       data={'applicability': value}, follow_redirects=True)


def text(resp):
    return resp.data.decode('utf-8')


# ── Bug A ─────────────────────────────────────────────────────────────

def test_a_primeros_auxilios_stays_deactivated(client):
    emp_id = new_emp('LB')
    client.get(u('employee_detail', id=emp_id))            # backfill creates slots, PA=YES
    (pa,) = emp_rows(emp_id, 'PRIMEROS_AUXILIOS')
    assert pa.applicability == 'YES'
    r = toggle_emp(client, emp_id, pa.id, 'N/A')
    assert 'desactivado' in text(r)
    client.get(u('employee_detail', id=emp_id))            # the page load that used to undo it
    (pa,) = emp_rows(emp_id, 'PRIMEROS_AUXILIOS')
    assert pa.applicability == 'N/A', 'toggle to N/A was undone by the detail-page backfill'


# ── Bug B: several "Otro" per record ──────────────────────────────────

def test_b_two_other_permits(client):
    emp_id = new_emp('LB')
    client.get(u('employee_detail', id=emp_id))
    r1 = add_emp_permit(client, emp_id, 'OTHER', 'Manejo de extintor de incendios')
    assert 'ya tiene el permiso' not in text(r1)
    r2 = add_emp_permit(client, emp_id, 'OTHER', 'Curso de montacargas')
    assert 'ya tiene el permiso' not in text(r2)
    rows = emp_rows(emp_id, 'OTHER')
    assert sorted(p.permit_name for p in rows) == ['Curso de montacargas', 'Manejo de extintor de incendios']
    assert all(p.applicability == 'YES' for p in rows)


def test_c_other_same_name_active_warns(client):
    emp_id = new_emp('LB')
    add_emp_permit(client, emp_id, 'OTHER', 'Manejo de extintor de incendios')
    r = add_emp_permit(client, emp_id, 'OTHER', '  manejo de EXTINTOR de incendios ')
    assert 'ya tiene el permiso' in text(r)
    assert len(emp_rows(emp_id, 'OTHER')) == 1


def test_d_other_na_reactivates_by_name_keeping_data(client):
    emp_id = new_emp('LB')
    add_emp_permit(client, emp_id, 'OTHER', 'Manejo de extintor de incendios')
    add_emp_permit(client, emp_id, 'OTHER', 'Curso de montacargas')
    with app.app_context():
        ext = EmployeePermit.query.filter_by(employee_id=emp_id, permit_name='Manejo de extintor de incendios').one()
        ext.permit_number = 'X-1'
        db.session.commit()
        ext_id = ext.id
    toggle_emp(client, emp_id, ext_id, 'N/A')
    r = add_emp_permit(client, emp_id, 'OTHER', 'Manejo de extintor de incendios', '2027-01-15')
    assert 'ya tiene el permiso' not in text(r)
    rows = emp_rows(emp_id, 'OTHER')
    assert len(rows) == 2
    ext = next(p for p in rows if p.id == ext_id)
    assert ext.applicability == 'YES'
    assert ext.expiration_date == date(2027, 1, 15)
    assert ext.permit_number == 'X-1'
    other = next(p for p in rows if p.id != ext_id)
    assert other.permit_name == 'Curso de montacargas' and other.applicability == 'YES'


def test_e_other_blank_name_rejected(client):
    emp_id = new_emp('LB')
    r = add_emp_permit(client, emp_id, 'OTHER', '   ')
    assert 'Indique el nombre' in text(r)
    assert emp_rows(emp_id, 'OTHER') == []


def test_f_active_standard_type_warns(client):
    emp_id = new_emp('LB')
    client.get(u('employee_detail', id=emp_id))
    (ntsp,) = emp_rows(emp_id, 'NTSP')
    toggle_emp(client, emp_id, ntsp.id, 'YES')
    r = add_emp_permit(client, emp_id, 'NTSP', '', '2027-03-01')
    assert 'ya tiene el permiso' in text(r)
    assert len(emp_rows(emp_id, 'NTSP')) == 1


def test_g_na_standard_type_reactivates(client):
    emp_id = new_emp('LB')
    client.get(u('employee_detail', id=emp_id))
    (twic,) = emp_rows(emp_id, 'TWIC')
    assert twic.applicability == 'N/A'
    r = add_emp_permit(client, emp_id, 'TWIC', '', '2027-03-01')
    assert 'agregado' in text(r)
    (twic,) = emp_rows(emp_id, 'TWIC')
    assert twic.applicability == 'YES' and twic.expiration_date == date(2027, 3, 1)


def test_h_equipment_add_and_toggle(client):
    eq_id = new_equipment('LB')
    add_eq_permit(client, eq_id, 'OTHER', 'Permiso DTOP')
    r = add_eq_permit(client, eq_id, 'OTHER', 'Permiso Municipal')
    assert 'ya tiene el permiso' not in text(r)
    assert len(eq_rows(eq_id, 'OTHER')) == 2
    # policy-covered vehicle: SEGURO add is refused, row stays N/A
    (seg,) = eq_rows(eq_id, 'SEGURO')
    assert seg.applicability == 'N/A'
    r = add_eq_permit(client, eq_id, 'SEGURO', '', '2027-01-01')
    assert 'cubierto por una p' in text(r)
    (seg,) = eq_rows(eq_id, 'SEGURO')
    assert seg.applicability == 'N/A'
    # standard slot: toggle off then re-add reactivates the same row
    (marb,) = eq_rows(eq_id, 'MARBETE')
    toggle_eq(client, eq_id, marb.id, 'N/A')
    r = add_eq_permit(client, eq_id, 'MARBETE', '', '2027-02-02')
    assert 'agregado' in text(r)
    rows = eq_rows(eq_id, 'MARBETE')
    assert len(rows) == 1 and rows[0].id == marb.id
    assert rows[0].applicability == 'YES' and rows[0].expiration_date == date(2027, 2, 2)
    # own-policy (Personal) vehicle: SEGURO is editable
    p_id = new_equipment('Personal')
    (seg,) = eq_rows(p_id, 'SEGURO')
    assert seg.applicability == 'YES'
    toggle_eq(client, p_id, seg.id, 'N/A')
    assert eq_rows(p_id, 'SEGURO')[0].applicability == 'N/A'
    r = add_eq_permit(client, p_id, 'SEGURO', '', '2027-05-05')
    assert 'agregado' in text(r)
    assert eq_rows(p_id, 'SEGURO')[0].applicability == 'YES'


def test_i_db_level_uniqueness(client):
    emp_id = new_emp('LB')
    client.get(u('employee_detail', id=emp_id))
    with app.app_context():
        db.session.add(EmployeePermit(employee_id=emp_id, permit_type='NTSP', applicability='YES'))
        with pytest.raises(IntegrityError):
            db.session.commit()
        db.session.rollback()
        for n in ('a', 'b', 'c'):
            db.session.add(EmployeePermit(employee_id=emp_id, permit_type='OTHER', permit_name=n, applicability='YES'))
        db.session.commit()
        assert EmployeePermit.query.filter_by(employee_id=emp_id, permit_type='OTHER').count() == 3


def test_j_postgres_ddl_and_dedup_lists():
    from sqlalchemy.schema import CreateIndex
    from sqlalchemy.dialects import postgresql
    pg = postgresql.dialect()
    expected = {
        'uq_employee_permit_type_partial': "CREATE UNIQUE INDEX uq_employee_permit_type_partial ON employee_permits "
                                           "(employee_id, permit_type) WHERE permit_type <> 'OTHER'",
        'uq_equipment_permit_type_partial': "CREATE UNIQUE INDEX uq_equipment_permit_type_partial ON equipment_permits "
                                            "(equipment_id, permit_type) WHERE permit_type <> 'OTHER'",
    }
    found = {}
    for table in (EmployeePermit.__table__, EquipmentPermit.__table__):
        for idx in table.indexes:
            if idx.name in expected:
                found[idx.name] = str(CreateIndex(idx).compile(dialect=pg)).strip()
        from sqlalchemy import UniqueConstraint
        for c in table.constraints:
            if isinstance(c, UniqueConstraint):
                assert {col.name for col in c.columns} != {'employee_id', 'permit_type'}
                assert {col.name for col in c.columns} != {'equipment_id', 'permit_type'}
    assert found == expected
    ddl = ' '.join(sql for _label, sql in appmod.PERMIT_UNIQUENESS_DDL)
    assert ddl.count('IF NOT EXISTS') == 2 and ddl.count('DROP CONSTRAINT IF EXISTS') == 2
    assert "WHERE permit_type <> 'OTHER'" in ddl
    src = pathlib.Path(appmod.__file__).read_text()
    assert 'uq_employee_permit_type UNIQUE' not in src
    assert 'uq_equipment_permit_type UNIQUE' not in src


# ── Blanket (company-wide) permit types ───────────────────────────────

CUSTOM_LABEL = 'Manejo de extintor'
CUSTOM_CODE = 'CUSTOM_MANEJO_DE_EXTINTOR'


def _custom_type():
    with app.app_context():
        return models.EmployeePermitType.query.filter_by(code=CUSTOM_CODE).first()


@pytest.fixture(scope='module')
def blanket(client):
    """Create the PLI-only blanket type once; returns the ids of the fixtures around it."""
    ids = {
        'pli': new_emp('PLI'), 'pli2': new_emp('PLI'), 'lb': new_emp('LB'),
        'pli_archived': new_emp('PLI', archived=True),
    }
    r = client.post(u('admin_permit_type_new'), data={'label': CUSTOM_LABEL, 'company': 'PLI'},
                    follow_redirects=True)
    assert r.status_code == 200
    assert 'creado' in text(r)
    ids['type_id'] = _custom_type().id
    return ids


def test_k_create_blanket_type_for_pli(client, blanket):
    for key in ('pli', 'pli2'):
        (row,) = emp_rows(blanket[key], CUSTOM_CODE)
        assert row.applicability == 'YES' and row.expiration_date is None
        assert row.status == 'missing'
        with app.app_context():
            assert EmployeePermit.query.get(row.id).display_name == CUSTOM_LABEL
    assert emp_rows(blanket['lb'], CUSTOM_CODE) == []
    assert emp_rows(blanket['pli_archived'], CUSTOM_CODE) == []
    r = client.post(u('admin_permit_type_new'), data={'label': CUSTOM_LABEL, 'company': 'PLI'},
                    follow_redirects=True)
    assert 'ya existe' in text(r)
    r = client.post(u('admin_permit_type_new'), data={'label': 'NTSP', 'company': 'LB'},
                    follow_redirects=True)
    assert 'reservado' in text(r)
    with app.app_context():
        assert models.EmployeePermitType.query.count() == 1
    page = text(client.get(u('admin_permit_types')))
    assert CUSTOM_LABEL in page and CUSTOM_CODE in page


def test_l_new_and_existing_employees_get_the_slot(client, blanket):
    r = client.post(u('employee_new'), data={'name': _uniq('Nuevo PLI'), 'company': 'PLI'},
                    follow_redirects=True)
    assert r.status_code == 200
    with app.app_context():
        emp = Employee.query.filter(Employee.name.like('Nuevo PLI%')).first()
    (row,) = emp_rows(emp.id, CUSTOM_CODE)
    assert row.applicability == 'YES'
    # LB employee page: select does not offer it, no card
    lb_page = text(client.get(u('employee_detail', id=blanket['lb'])))
    assert CUSTOM_CODE not in lb_page
    # PLI employee page: card present, shows Incompleto
    pli_page = text(client.get(u('employee_detail', id=blanket['pli'])))
    assert CUSTOM_LABEL in pli_page and 'Incompleto' in pli_page
    # moving the LB employee to PLI adds the slot
    client.post(u('employee_edit', id=blanket['lb']), data={'name': 'Moved To PLI', 'company': 'PLI'},
                follow_redirects=True)
    (row,) = emp_rows(blanket['lb'], CUSTOM_CODE)
    assert row.applicability == 'YES'
    # and moving back deactivates it (data kept) rather than deleting it
    client.post(u('employee_edit', id=blanket['lb']), data={'name': 'Moved To PLI', 'company': 'LB'},
                follow_redirects=True)
    (row,) = emp_rows(blanket['lb'], CUSTOM_CODE)
    assert row.applicability == 'N/A'


def test_m_dashboard_filter_is_company_aware(client, blanket):
    with app.app_context():
        lb = Employee.query.get(blanket['lb'])
        pli = Employee.query.get(blanket['pli'])
        assert appmod._permit_type_status(lb, CUSTOM_CODE, 'employees') == 'na'
        assert appmod._permit_type_status(pli, CUSTOM_CODE, 'employees') == 'missing'
    page = text(client.get(f'/?view=employees&permit_type={CUSTOM_CODE}'))
    assert f'value="{CUSTOM_CODE}"' in page   # option survived validation
    with app.app_context():
        ctx = appmod._dashboard_context({'view': 'employees', 'permit_type': CUSTOM_CODE})
        assert ctx['permit_type'] == CUSTOM_CODE
        names = [e.name for e in ctx['lb_items']] + [e.name for e in ctx['pli_items']]
        assert all(Employee.query.filter_by(name=n).first().company == 'PLI' for n in names)
        assert names  # the PLI employees with the missing date are flagged
    client.get('/?reset=1')


def test_n_excel_and_pdf_columns(client, blanket):
    from openpyxl import load_workbook
    r = client.get('/export/excel?type=employees')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb['Empleados']
    headers = [c.value for c in ws[1]]
    assert CUSTOM_LABEL in headers
    col = headers.index(CUSTOM_LABEL) + 1
    by_company = {}
    for row in ws.iter_rows(min_row=2):
        by_company.setdefault(row[0].value, set()).add(row[col - 1].value)
    assert by_company['LB'] == {'N/A'}
    assert '—' in by_company['PLI']           # active slot with no date reads as missing
    r = client.get('/report/pdf?type=employees')
    html = text(r)
    lb_block = html.split('LB Caribe Services — Empleados')[1].split('Professional Logistics — Empleados')[0]
    pli_block = html.split('Professional Logistics — Empleados')[1].split('</table>')[0]
    assert CUSTOM_LABEL not in lb_block
    assert CUSTOM_LABEL in pli_block
    for block in (lb_block, pli_block):
        th = block.count('<th>')
        first_row = block.split('<tbody>')[1].split('</tr>')[0]
        assert first_row.count('<td') == th, 'PDF header/cell count drifted'


def test_o_notifications_pick_up_custom_type(client, blanket):
    with app.app_context():
        emp = Employee.query.get(blanket['pli'])
        emp.email = 'pli@example.com'
        row = EmployeePermit.query.filter_by(employee_id=emp.id, permit_type=CUSTOM_CODE).one()
        row.expiration_date = date.today() - timedelta(days=3)
        db.session.commit()
        items = appmod._gather_expiring_items(employee_id=emp.id)
        assert CUSTOM_LABEL in str(items), items


def test_p_delete_type_removes_rows_and_files(client, blanket):
    with app.app_context():
        row = EmployeePermit.query.filter_by(employee_id=blanket['pli2'], permit_type=CUSTOM_CODE).one()
        db.session.add(FileStorage(filename='unit_test_doc.pdf', original_filename='doc.pdf', data=b'%PDF-1.4'))
        row.file_path = 'unit_test_doc.pdf'
        db.session.commit()
        other_count = EmployeePermit.query.filter(EmployeePermit.permit_type != CUSTOM_CODE).count()
    r = client.post(u('admin_permit_type_delete', id=blanket['type_id']), follow_redirects=True)
    assert r.status_code == 200
    with app.app_context():
        assert models.EmployeePermitType.query.filter_by(code=CUSTOM_CODE).first() is None
        assert EmployeePermit.query.filter_by(permit_type=CUSTOM_CODE).count() == 0
        assert FileStorage.query.filter_by(filename='unit_test_doc.pdf').first() is None
        assert EmployeePermit.query.count() == other_count
        ctx = appmod._dashboard_context({'view': 'employees', 'permit_type': CUSTOM_CODE})
        assert ctx['permit_type'] == ''


def test_q_dedup_cli_runs_and_matches_other_by_name(client):
    """`flask dedup` must still be the decorated command, and its merge helper
    pairs OTHER permits by normalized name (unmatched ones are re-parented)."""
    from click.testing import CliRunner
    with app.app_context():
        a = Employee(name=_uniq('Keeper'), company='LB')
        b = Employee(name=_uniq('Dup'), company='LB')
        db.session.add_all([a, b])
        db.session.flush()
        db.session.add(EmployeePermit(employee_id=a.id, permit_type='OTHER', permit_name='Extintor', applicability='YES'))
        pb1 = EmployeePermit(employee_id=b.id, permit_type='OTHER', permit_name='  extintor ', applicability='YES')
        pb2 = EmployeePermit(employee_id=b.id, permit_type='OTHER', permit_name='Montacargas', applicability='YES')
        pb3 = EmployeePermit(employee_id=b.id, permit_type='NTSP', applicability='YES')
        db.session.add_all([pb1, pb2, pb3])
        db.session.commit()
        assert appmod._matching_keeper_permit(a, pb1).permit_name == 'Extintor'
        assert appmod._matching_keeper_permit(a, pb2) is None
        assert appmod._matching_keeper_permit(a, pb3) is None   # keeper has no NTSP row yet
        for args in (['dedup', '--dry-run'], ['dedup']):
            r = CliRunner().invoke(app.cli, args)
            assert r.exit_code == 0, r.output or r.exception
            assert 'Dedup complete' in r.output
